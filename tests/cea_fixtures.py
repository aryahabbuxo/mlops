"""Builders for tiny synthetic CEA workbooks that mimic the real layout."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook

KARNATAKA_ROW = 28
LAST_ROW = 47
DEFAULT_SHEETS = ("Generation", "ISGS", "State-Wise")


def report_text(day: date) -> str:
    """Date text as written in cell B5, e.g. 06-July-24."""
    return day.strftime("%d-%B-%y")


def write_report(
    path: Path,
    day: date,
    *,
    all_values: tuple[float, float] = (10.0, 20.0),
    state_values: tuple[float, float] = (8.0, 15.0),
    sheet_order: tuple[str, ...] = DEFAULT_SHEETS,
    sheet_dates: dict[str, date] | None = None,
    karnataka_rows: int = 1,
    india_values: tuple[float, float] = (100.0, 200.0),
) -> Path:
    """Write a workbook with the sheets, B5 date and Karnataka row the parser expects."""
    sheet_dates = sheet_dates or {}
    wb = Workbook()
    wb.remove(wb.active)
    for name in sheet_order:
        ws = wb.create_sheet(name)
        if name not in ("Generation", "State-Wise"):
            ws["A1"] = "ISGS placeholder"
            continue
        wind, solar = all_values if name == "Generation" else state_values
        ws["A3"] = "Daily Renewable Generation Report (All India ) / title"
        ws["B5"] = report_text(sheet_dates.get(name, day))
        for row in range(7, LAST_ROW + 1):
            ws.cell(row=row, column=1, value=f"state {row}")
        ws.cell(row=KARNATAKA_ROW, column=1, value="कर्नाटक / Karnataka")
        ws.cell(row=KARNATAKA_ROW, column=2, value=wind)
        ws.cell(row=KARNATAKA_ROW, column=3, value=solar)
        ws.cell(row=LAST_ROW, column=1, value="सम्पूर्ण भारत / All India")
        ws.cell(row=LAST_ROW, column=2, value=india_values[0])
        ws.cell(row=LAST_ROW, column=3, value=india_values[1])
        for extra in range(1, karnataka_rows):
            ws.cell(row=KARNATAKA_ROW + extra, column=1, value="Karnataka again")
    wb.save(path)
    return path
